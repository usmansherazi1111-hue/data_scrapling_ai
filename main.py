from __future__ import annotations
from urllib.parse import urlparse
import asyncio, csv, io, json, os, re, time, uuid
from fastapi import FastAPI, HTTPException, Request, Form
from fastapi.responses import FileResponse, StreamingResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from dotenv import load_dotenv

load_dotenv()

from scraper.db import HistoryDB
from scraper.engine import CrawlEngine
from scraper.urltools import normalize_url
from scraper.auth import ProfileStore, parse_cookies
from scraper.enrichment import flatten
from scraper import ai, security
from scraper import opencorporates as oc, registries, overture, gapfill
from scraper.models import now_iso

VERSION = "2.1.1"
app = FastAPI(title="Scrapling Studio", version=VERSION)
app.mount("/static", StaticFiles(directory="static"), name="static")
DB_PATH = os.getenv("DB_PATH", "./data/scrapling.db")
DB = HistoryDB(DB_PATH); DB.fail_stale(); JOBS = {}; TASKS = set(); REGISTRY_RUNS = {}
PROFILES = ProfileStore(os.path.join(os.path.dirname(DB_PATH) or ".", "profiles"))
ENGINE = CrawlEngine(DB, JOBS, PROFILES)


def background(coro):
    task = asyncio.create_task(coro)
    TASKS.add(task); task.add_done_callback(TASKS.discard)


class CrawlRequest(BaseModel):
    url: str
    mode: str = Field(default="auto", pattern="^(auto|http|dynamic|stealth)$")
    max_pages: int = Field(default=20, ge=1, le=500)
    depth: int = Field(default=2, ge=0, le=5)
    concurrency: int = Field(default=4, ge=1, le=16)
    browser_pages: int = Field(default=3, ge=1, le=4)
    adaptive: bool = True
    network_idle: bool = False
    disable_resources: bool = True
    capture_xhr: bool = False
    ai_enrichment: bool = True
    registry_lookup: bool = True
    enrich: bool = True
    escalate: bool = True
    respect_robots: bool = True
    delay_ms: int = Field(default=0, ge=0, le=60000)
    proxies: str = ""
    profile: str | None = None
    timeout: int = Field(default=30000, ge=5000, le=120000)


@app.middleware("http")
async def require_login(request: Request, call_next):
    open_paths = ("/login", "/static/app.css")
    if request.url.path in open_paths or security.is_authenticated(request.cookies.get(security.COOKIE)):
        response = await call_next(request)
        if request.url.path == "/" or request.url.path.startswith("/static/"):
            response.headers["Cache-Control"] = "no-cache"  # always pick up a new version of the UI after an update
        return response
    if request.url.path.startswith("/api/"): return JSONResponse({"detail": "Not signed in"}, status_code=401)
    return RedirectResponse("/login", status_code=303)


@app.get("/login")
async def login_page(): return HTMLResponse(security.LOGIN_PAGE.format(error=""))


@app.post("/login")
async def login(request: Request, password: str = Form("")):
    if not security.check_password(password):
        await asyncio.sleep(1)  # slow down guessing
        return HTMLResponse(security.LOGIN_PAGE.format(error='<p class="err small">Wrong password.</p>'), status_code=401)
    r = RedirectResponse("/", status_code=303)
    r.set_cookie(security.COOKIE, security.session_token(), httponly=True, samesite="lax", max_age=30 * 86400,
                 secure=request.headers.get("x-forwarded-proto") == "https" or request.url.scheme == "https")
    return r


@app.get("/logout")
async def logout():
    r = RedirectResponse("/login", status_code=303); r.delete_cookie(security.COOKIE); return r


@app.get("/")
async def index(): return FileResponse("static/index.html")


@app.get("/api/status")
async def status():
    return {"version": VERSION, "ai": ai.configured(), "ai_model": os.getenv("AI_MODEL", "gpt-4.1-mini") if ai.configured() else None, "hunter": bool(os.getenv("HUNTER_API_KEY")), "opencorporates": oc.configured(), "registries": registries.available()}


@app.post("/api/jobs")
async def create_job(req: CrawlRequest):
    url = normalize_url(req.url)
    if not url or not url.startswith(("http://", "https://")): raise HTTPException(400, "Invalid URL")
    if not await security.is_public_url(url): raise HTTPException(400, "Private, local and internal addresses cannot be crawled")
    if req.profile and not PROFILES.get(req.profile): raise HTTPException(400, f"Auth profile '{req.profile}' not found")
    job_id = uuid.uuid4().hex[:12]
    background(ENGINE.run(job_id, url, req.model_dump()))
    return {"job_id": job_id}


def result_for(job_id):
    r = JOBS.get(job_id) or DB.get(job_id)
    if not r: raise HTTPException(404, "Job not found")
    return r.jsonable() if hasattr(r, "jsonable") else r


@app.get("/api/jobs/{job_id}")
async def job(job_id: str): return result_for(job_id)


@app.post("/api/jobs/{job_id}/enrich")
async def enrich(job_id: str, ai_enrichment: bool = True, registry_lookup: bool = True):
    r = result_for(job_id)
    if r["status"] == "running": raise HTTPException(409, "Crawl still running")
    if (r.get("enrichment") or {}).get("status") == "running": raise HTTPException(409, "Enrichment already running")
    background(ENGINE.enrich(job_id, use_ai=ai_enrichment, use_registry=registry_lookup))
    return {"ok": True}


@app.get("/api/history")
async def history(): return DB.list()


# ---------------------------------------------------------------- auth profiles
class ProfileRequest(BaseModel):
    name: str
    domain: str = ""
    cookies: str = ""
    headers: str = ""
    useragent: str = ""


@app.get("/api/profiles")
async def profiles(): return PROFILES.list()


@app.post("/api/profiles")
async def save_profile(req: ProfileRequest):
    try:
        cookies = parse_cookies(req.cookies, req.domain) if req.cookies.strip() else None
        headers = None
        if req.headers.strip():
            headers = {k.strip(): v.strip() for k, _, v in (l.partition(":") for l in req.headers.splitlines()) if k.strip() and v.strip()}
        PROFILES.save(req.name, cookies=cookies, headers=headers, useragent=req.useragent or None, domain=req.domain or None)
    except (ValueError, json.JSONDecodeError) as e:
        raise HTTPException(400, str(e))
    return {"ok": True}


@app.delete("/api/profiles/{name}")
async def delete_profile(name: str):
    try:
        await PROFILES.close_browser(name); PROFILES.delete(name)
    except ValueError as e: raise HTTPException(400, str(e))
    return {"ok": True}


class BrowserRequest(BaseModel):
    url: str


@app.post("/api/profiles/{name}/browser")
async def open_browser(name: str, req: BrowserRequest, request: Request):
    if security.is_remote(request.headers):
        raise HTTPException(403, "'Open browser' only works on the computer running the app (the window opens there). Remote users can import cookies instead.")
    url = normalize_url(req.url)
    if not url.startswith(("http://", "https://")): raise HTTPException(400, "Invalid URL")
    try: await PROFILES.open_browser(name, url)
    except ValueError as e: raise HTTPException(400, str(e))
    except Exception as e: raise HTTPException(500, f"Could not open browser: {e}")
    return {"ok": True}


@app.post("/api/profiles/{name}/browser/finish")
async def finish_browser(name: str):
    await PROFILES.close_browser(name)
    return {"ok": True, "profile": next((p for p in PROFILES.list() if p["name"] == name), None)}


# ---------------------------------------------------------------- company lists: industry + country -> companies -> automatic website crawl -> Excel
MAX_LIST = int(os.getenv("COMPANY_LIST_MAX", "200"))
CRAWL_PAGES = int(os.getenv("COMPANY_CRAWL_PAGES", "6"))
# Latency caps for list crawls: most sites answer in a few seconds; a few protected ones take minutes and would stall the list.
CRAWL_FETCH_MS = int(os.getenv("COMPANY_FETCH_TIMEOUT_MS", "15000"))
CRAWL_PAGE_SECONDS = float(os.getenv("COMPANY_PAGE_SECONDS", "25"))
CRAWL_SITE_SECONDS = float(os.getenv("COMPANY_SITE_SECONDS", "45"))
CRAWL_CONCURRENCY = int(os.getenv("COMPANY_CRAWL_CONCURRENCY", "8"))


class CompanySearchRequest(BaseModel):
    industry: str = Field(default="", max_length=100)   # any words: "textile", "medicine", "software"
    country: str = Field(default="", max_length=60)     # "pk", "Pakistan", "UK"
    name: str = Field(default="", max_length=100)       # optional: only companies whose name contains this
    limit: int = Field(default=50, ge=5, le=MAX_LIST)   # how many companies to keep and crawl
    category: str = Field(default="", max_length=100)   # optional: only this category (the chips under the results)
    crawl: bool = True


def new_run(label, query):
    run = {"id": uuid.uuid4().hex[:12], "created_at": now_iso(), "kind": "companies", "label": label[:200], "status": "running", "query": query, "source": "overture",
           "companies": [], "progress": "starting", "error": None, "notes": [], "total_count": None, "crawl": {"status": "idle", "done": 0, "total": 0, "failed": 0}}
    REGISTRY_RUNS[run["id"]] = run; DB.save_registry(run); return run


def company_site(c: dict) -> str | None:
    """The company's home page (Overture often stores a deep link such as /careers)."""
    w = (c.get("website") or "").strip()
    if not w: return None
    u = normalize_url(w if w.startswith(("http://", "https://")) else "http://" + w)
    return re.sub(r"^(https?://[^/]+).*$", r"\1", u) if u else None


def keep_profile(enrichment: dict | None) -> dict:
    e = enrichment or {}
    return {"people": e.get("people", [])[:10], "emails": (e.get("contacts") or {}).get("emails", [])[:30], "phones": (e.get("contacts") or {}).get("phones", [])[:15],
            "social": {k: v.get("url") for k, v in (e.get("social") or {}).items()}, "tech_stack": [t.get("name") for t in e.get("tech_stack", [])][:30]}


async def domain_resolves(url: str) -> bool:
    """Dead domains are common in map data; failing them here saves launching a browser for each one."""
    import socket
    host = urlparse(url).hostname or ""
    try:
        await asyncio.wait_for(asyncio.get_running_loop().getaddrinfo(host, None, type=socket.SOCK_STREAM), 8)
        return True
    except socket.gaierror:
        return False
    except Exception:
        return True   # slow or odd resolver: let the crawl decide


def maybe_save(run, force=False):
    """Writing the whole list after every company is O(n squared); save at most every few seconds."""
    now = time.monotonic()
    if force or now - run.get("_saved", 0) > 3:
        run["_saved"] = now; DB.save_registry({k: v for k, v in run.items() if not k.startswith("_")})


async def crawl_site(c: dict, url: str):
    job_id = uuid.uuid4().hex[:12]; c["job_id"] = job_id
    cfg = CrawlRequest(url=url, max_pages=CRAWL_PAGES, depth=1, ai_enrichment=True, registry_lookup=False, timeout=CRAWL_FETCH_MS, concurrency=CRAWL_PAGES).model_dump()
    cfg.update(page_time_limit=CRAWL_PAGE_SECONDS, time_budget=CRAWL_SITE_SECONDS, markdown=False, parse_in_thread=os.getenv("COMPANY_PARSE_IN_THREAD", "1") != "0")
    res = (await ENGINE.run(job_id, url, cfg)).jsonable()
    readable = "pages" not in res or any(200 <= (p.get("status") or 0) < 400 for p in res["pages"])
    return res, readable


async def process_company(c: dict, cc: str, sem, state, run):
    """Listing -> (missing website? ask OpenStreetMap) -> crawl site (alternate address if it will not open) -> official register -> best value per field."""
    try:
        await gapfill.discover_listing(c, cc)
        registry = asyncio.create_task(gapfill.registry_fill(c, cc))   # runs alongside the crawl
        async with sem:
            url = company_site(c)
            c["crawl_status"] = "running"; state["current"] = c.get("name")
            if not url:
                c["crawl_status"] = "no website"
            else:
                tried, error = [], None
                for cand in [url] + gapfill.alt_urls(url):
                    try:
                        if not await security.is_public_url(cand): error = "website is not a public address"; continue
                        if not await domain_resolves(cand): error = "website domain does not exist (DNS)"; continue
                        res, readable = await crawl_site(c, cand)
                    except Exception as e:
                        error = str(e)[:200]; tried.append(cand); continue
                    tried.append(cand)
                    if readable:
                        c["crawl_status"] = res["status"]; error = None
                        c["details"] = {k: v for k, v in flatten(res.get("enrichment")).items() if v not in (None, "", False)}
                        c["profile"] = keep_profile(res.get("enrichment"))
                        if cand != url: c["website"] = cand; c.setdefault("filled_from", {})["website"] = "alternate address"
                        break
                    codes = sorted({str(p.get("status") or p.get("error") or "no answer")[:40] for p in res.get("pages", [])})
                    error = "website not readable (" + ", ".join(codes[:3]) + ")"
                    if os.getenv("COMPANY_TRY_ALT_URLS", "1") == "0": break
                if error:
                    c["crawl_status"] = "failed"; c["crawl_error"] = error; state["failed"] += 1
        try: await asyncio.wait_for(registry, 60)
        except Exception: registry.cancel()
    except Exception as e:
        c["crawl_status"] = "failed"; c["crawl_error"] = str(e)[:200]; state["failed"] += 1
    gapfill.merge_best(c)
    state["done"] += 1; maybe_save(run)


async def run_company_search(run, req: CompanySearchRequest):
    error = None; cc = overture.country_code(req.country)
    try:
        if cc:
            run["progress"] = f"reading open company data for {overture.country_name(cc)}"
            def progress(m): run["progress"] = m
            res = await overture.search(req.industry, req.country, req.name, req.limit, category=req.category, progress=progress)
            run["companies"] = res["companies"]; run["total_count"] = res["matched"]; run["facets"] = res.get("facets", [])
            n_weak = sum(1 for c in res["companies"] if c.get("match") == "name only")
            run["notes"].append(f"{res['matched']:,} companies match ({res.get('relevant', 0):,} by category). Showing the best {len(res['companies'])}: category matches first, then those with the most contact details.")
            if n_weak: run["notes"].append(f"{n_weak} are listed because only their name contains the word (check the Why listed column), or narrow by category below.")
            if res.get("partial"): run["notes"].append(res["partial"])
        elif req.country.strip():
            raise overture.OvertureError(f"Unknown country '{req.country.strip()}'. Use a country name or a 2-letter code like pk, gb or us.")
        elif req.name.strip():   # no country: look the name up in the official LEI register
            run["source"] = "gleif"; run["progress"] = "searching GLEIF by name"
            res = await registries.search("gleif", req.name.strip(), per_page=min(req.limit, 100))
            run["companies"] = res["companies"]; run["total_count"] = res["total_count"]
            run["notes"].append("No country given, so the official LEI register (GLEIF) was searched by name. Add a country and an industry to list companies.")
        else:
            raise overture.OvertureError("Enter an industry and a country.")
    except Exception as e:
        error = str(e)[:500]
    run["status"] = "failed" if error and not run["companies"] else "completed"; run["error"] = error
    maybe_save(run, True)
    if req.crawl and run["companies"]:
        todo = run["companies"] if cc else [c for c in run["companies"] if company_site(c)]
        state = run["crawl"] = {"status": "running", "done": 0, "total": len(todo), "failed": 0, "current": None, "started_at": now_iso()}
        run["progress"] = f"collecting contact details for {len(todo)} companies"; maybe_save(run, True)
        sem = asyncio.Semaphore(CRAWL_CONCURRENCY)
        try:
            await asyncio.gather(*(process_company(c, cc or "", sem, state, run) for c in todo))
        finally:
            state.update(status="completed", current=None, finished_at=now_iso())
    run["progress"] = "done"; run["finished_at"] = now_iso(); maybe_save(run, True); REGISTRY_RUNS.pop(run["id"], None)


@app.post("/api/companies/search")
async def companies_search(req: CompanySearchRequest):
    if not (req.industry.strip() or req.name.strip()): raise HTTPException(400, "Enter an industry (for example textile) and a country (for example pk)")
    if req.industry.strip() and not req.country.strip(): raise HTTPException(400, "Enter a country too, for example pk or Pakistan")
    run = new_run(" · ".join(x for x in (req.industry.strip(), req.country.strip(), req.name.strip(), req.category.strip()) if x), req.model_dump())
    background(run_company_search(run, req))
    return {"id": run["id"]}


@app.get("/api/companies")
async def companies_lists(): return DB.list_registry()


def run_for(list_id):
    r = REGISTRY_RUNS.get(list_id) or DB.get_registry(list_id)
    if not r: raise HTTPException(404, "List not found")
    return r


@app.get("/api/companies/{list_id}")
async def companies_list(list_id: str): return run_for(list_id)


LICENCES = {"overture": "Overture Maps places: (c) Overture Maps Foundation, CDLA Permissive 2.0. https://overturemaps.org",
            "gleif": "GLEIF LEI data: CC0 (no restrictions). https://www.gleif.org"}
SUMMARY_COLS = ["name", "category", "match", "city", "country_code", "website", "website_source", "email", "email_source", "phone", "phone_source", "decision_maker",
                "decision_maker_title", "decision_maker_email", "all_emails", "phones_e164", "social_linkedin", "social_facebook", "social_instagram", "social_twitter",
                "lead_score", "lead_grade", "tech_stack", "legal_name", "lei", "registry_status", "registered_address", "locations", "missing", "crawl_status", "crawl_error", "source"]
HIDE = {"source_id", "category_path", "confidence", "job_id", "jurisdiction_code", "filled_from", "best", "best_source", "registry", "profile", "details", "socials"}


def summary_row(c):
    d = dict(c.get("details") or {}); b = c.get("best") or {}; bs = c.get("best_source") or {}; g = c.get("registry") or {}
    row = {**{k: v for k, v in c.items() if k not in HIDE}, **d}
    for f in ("email", "phone", "website"):
        if b.get(f): row[f] = b[f]
        if bs.get(f): row[f + "_source"] = bs[f]
    for k, v in b.items():
        if k.startswith("social_"): row[k] = v
    if b.get("contact_person") and not row.get("decision_maker"): row["decision_maker"] = b["contact_person"]
    emails = [e.strip() for e in str(row.get("all_emails") or "").split(",") if e.strip()]
    if b.get("email") and b["email"] not in emails: emails.insert(0, b["email"])
    if emails: row["all_emails"] = ", ".join(emails)
    row.update(legal_name=g.get("name"), lei=g.get("lei") or g.get("company_number"), registry_status=g.get("current_status"))
    if g.get("registered_address") and not row.get("registered_address"): row["registered_address"] = g["registered_address"]
    row["missing"] = ", ".join(c.get("missing") or [])
    return {k: v for k, v in row.items() if not isinstance(v, dict)}


def sheet_name(i, name, used):
    base = re.sub(r"[\[\]:*?/\\]", " ", f"{i:03d} {name or 'company'}").strip()[:28].strip()
    n, k = base, 2
    while n.lower() in used: n = f"{base[:25]}~{k}"; k += 1
    used.add(n.lower()); return n


def build_workbook(r):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    cs = r["companies"]; wb = Workbook(); ws = wb.active; ws.title = "Summary"; used = {"summary", "sources"}
    extra = [k for k in dict.fromkeys(k for c in cs for k in summary_row(c)) if k not in SUMMARY_COLS and k not in HIDE]
    cols = [c for c in SUMMARY_COLS if any(summary_row(x).get(c) not in (None, "") for x in cs)] + [k for k in extra if any(summary_row(x).get(k) not in (None, "") for x in cs)]
    names = [sheet_name(i, c.get("name"), used) for i, c in enumerate(cs, 1)]
    ws.append(["sheet"] + cols)
    for n, c in zip(names, cs):
        row = summary_row(c); ws.append([n] + [(", ".join(map(str, v)) if isinstance(v := row.get(k), list) else v) for k in cols])
        cell = ws.cell(row=ws.max_row, column=1); cell.hyperlink = f"#'{n}'!A1"; cell.font = Font(color="0563C1", underline="single")
    for i, n in enumerate(names):
        c = cs[i]; s = wb.create_sheet(n); p = c.get("profile") or {}; row = summary_row(c)
        s.append(["Company", c.get("name")])
        for k, v in row.items():
            if k in ("name", "source_id", "job_id") or v in (None, "", []) or isinstance(v, (dict, list)): continue
            s.append([k.replace("_", " "), v])
        if c.get("crawl_error"): s.append(["crawl note", c["crawl_error"]])
        for title, hdr, items in (("People", ["name", "title", "email", "email_guess", "source_url"], p.get("people")),
                                  ("Emails", ["email", "type", "on_company_domain", "domain_accepts_mail"], p.get("emails")),
                                  ("Phones", ["raw", "e164", "valid", "country", "type"], p.get("phones"))):
            if not items: continue
            s.append([]); s.append([title]); s.cell(row=s.max_row, column=1).font = Font(bold=True, size=12); s.append(hdr)
            for cell in s[s.max_row]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="EEEEEE")
            for it in items: s.append([json.dumps(it.get(h), ensure_ascii=False) if isinstance(it.get(h), (list, dict)) else it.get(h) for h in hdr])
        s.append([]); s.append(["Back to summary"]); s.cell(row=s.max_row, column=1).hyperlink = "#'Summary'!A1"
        s.column_dimensions["A"].width = 28; s.column_dimensions["B"].width = 60
        for cell in s["A"][:1]: cell.font = Font(bold=True)
    src = wb.create_sheet("Sources")
    src.append([f"Search: {r['label']}"]); src.append([f"Created: {r['created_at']}"])
    for s_ in sorted({c.get("source") for c in cs if c.get("source")}): src.append([LICENCES.get(s_, s_)])
    for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for j, col in enumerate(ws.columns, 1): ws.column_dimensions[get_column_letter(j)].width = min(45, max(12, max(len(str(c.value or "")) for c in col[:60]) + 2))
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    return wb


@app.get("/api/companies/{list_id}/export/{kind}")
async def companies_export(list_id: str, kind: str):
    r = run_for(list_id); fn = f"companies_{re.sub(r'[^a-z0-9]+', '_', r['label'].lower()).strip('_')[:40]}_{list_id[:6]}"
    if kind == "json":
        return StreamingResponse(io.BytesIO(json.dumps(r, ensure_ascii=False, indent=2).encode()), media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{fn}.json"'})
    if kind == "csv":
        rows = [summary_row(c) for c in r["companies"]]
        cols = [k for k in SUMMARY_COLS if any(x.get(k) not in (None, "") for x in rows)]
        b = io.StringIO(); w = csv.writer(b); w.writerow(cols)
        for x in rows: w.writerow([", ".join(map(str, v)) if isinstance(v := x.get(k), list) else v for k in cols])
        return StreamingResponse(iter([b.getvalue().encode("utf-8-sig")]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{fn}.csv"'})
    if kind == "xlsx":
        out = io.BytesIO(); build_workbook(r).save(out); out.seek(0)
        return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fn}.xlsx"'})
    raise HTTPException(400, "Unsupported export")


# ---------------------------------------------------------------- exports
def export_row(r):
    row = {k: (v.get("value") if isinstance(v, dict) else v) for k, v in r["fields"].items()}
    row = {k: (None if v == "Not Found" else v) for k, v in row.items()}
    row.update(flatten(r.get("enrichment")))
    return {k: (json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else ", ".join(map(str, v)) if isinstance(v, list) else v) for k, v in row.items()}


@app.get("/api/jobs/{job_id}/export/{kind}")
async def export(job_id: str, kind: str):
    r = result_for(job_id); row = export_row(r); e = r.get("enrichment") or {}
    if kind == "json":
        return StreamingResponse(io.BytesIO(json.dumps(r, ensure_ascii=False, indent=2).encode()), media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{job_id}.json"'})
    if kind == "csv":
        b = io.StringIO(); w = csv.DictWriter(b, fieldnames=row.keys()); w.writeheader(); w.writerow(row)
        return StreamingResponse(iter([b.getvalue().encode("utf-8-sig")]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{job_id}.csv"'})
    if kind == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook(); ws = wb.active; ws.title = "Company"
        for k, v in row.items(): ws.append([k.replace("_", " "), v])
        sheets = {
            "People": (["name", "title", "email", "email_guess", "email_guess_confidence", "source", "source_url"], e.get("people", [])),
            "Emails": (["email", "type", "on_company_domain", "domain_accepts_mail", "free_provider", "disposable"], (e.get("contacts") or {}).get("emails", [])),
            "Phones": (["raw", "e164", "international", "valid", "country", "location", "type"], (e.get("contacts") or {}).get("phones", [])),
            "Tech stack": (["name", "category", "pages", "evidence", "source_url"], e.get("tech_stack", [])),
            "Pages": (["url", "category", "status", "mode", "elapsed_ms", "challenge", "error"], r.get("pages", [])),
        }
        if e.get("registry"):
            sheets["Registry officers"] = (["name", "position", "start_date", "end_date", "current", "opencorporates_url"], e["registry"].get("officers", []))
        for title, (cols, items) in sheets.items():
            s = wb.create_sheet(title); s.append(cols)
            for it in items: s.append([json.dumps(it.get(c), ensure_ascii=False) if isinstance(it.get(c), (list, dict)) else it.get(c) for c in cols])
        for s in wb.worksheets:
            for cell in s[1]: cell.font = Font(bold=True)
            for col in s.columns: s.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
        out = io.BytesIO(); wb.save(out); out.seek(0)
        return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{job_id}.xlsx"'})
    raise HTTPException(400, "Unsupported export")


if __name__ == "__main__":
    import socket, sys, uvicorn
    host, port = os.getenv("HOST", "127.0.0.1"), int(os.getenv("PORT", "8000"))
    with socket.socket() as s:
        if s.connect_ex(("127.0.0.1" if host in ("0.0.0.0", "") else host, port)) == 0:
            print(f"\nPort {port} is already in use, most likely by an older copy of Scrapling Studio that is still running.\n"
                  f"Close that window (or end its python.exe in Task Manager) and run this again, or set PORT=8001 in .env.\n")
            sys.exit(1)
    print(f"\nScrapling Studio {VERSION}  on http://{host}:{port}\n")
    if sys.platform == "win32": asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())  # Playwright needs subprocess support
    uvicorn.run(app, host=host, port=port, loop="asyncio")
