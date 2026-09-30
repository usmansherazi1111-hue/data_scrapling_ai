"""Industry + country -> companies (Overture rows, real sample data) -> automatic crawl -> Excel with one sheet per company."""
import io, importlib, pathlib, time
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from scraper import overture

SAMPLE = str(pathlib.Path(__file__).parent / "fixtures" / "places_sample.parquet")


def scan(industry, country="pk", **kw):
    return overture.scan(industry, country, dataset_path=SAMPLE, release="test", **kw)


def test_country_lookup():
    assert overture.country_code("pk") == "PK" and overture.country_code("Pakistan") == "PK" and overture.country_code("UK") == "GB"
    assert overture.country_code("Nowhereland") is None and overture.country_code("") is None
    assert len(overture.COUNTRIES) > 200


def test_any_industry_word_matches_by_category_or_name_only_in_the_country():
    r = scan("textile", limit=500)
    assert r["matched"] > 300
    assert all(c["country_code"] == "PK" for c in r["companies"])     # the India rows in the file are excluded
    assert any(c["category"] == "textile manufacturer" for c in r["companies"])
    top = r["companies"][:20]
    assert all(c["website"] for c in top)                              # best contact details come first
    assert scan("Textiles", limit=500)["matched"] == r["matched"]      # plural and case do not matter


def test_no_per_industry_special_cases_unknown_word_gives_empty_not_error():
    assert scan("zzzqqq")["companies"] == []


def test_name_filter_and_limit_and_dedupe():
    r = scan("textile", name="mills", limit=5)
    assert len(r["companies"]) <= 5 and all("mills" in c["name"].lower() for c in r["companies"])
    names = [(c["name"].lower(), (c["city"] or "").lower()) for c in scan("textile", limit=500)["companies"]]
    assert len(names) == len(set(names))


def test_errors_are_plain():
    with pytest.raises(overture.OvertureError, match="Unknown country"): scan("textile", "Nowhereland")
    with pytest.raises(overture.OvertureError, match="industry"): scan("", "pk")


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_PATH", str(tmp_path / "t.db")); monkeypatch.setenv("APP_PASSWORD", "")
    import main; importlib.reload(main)
    orig = overture.scan
    monkeypatch.setattr(overture, "scan", lambda *a, **k: orig(*a, dataset_path=SAMPLE, release="test", **{x: y for x, y in k.items() if x != "release"}))
    with TestClient(main.app) as c:
        yield c


def wait(c, rid):
    for _ in range(200):
        r = c.get(f"/api/companies/{rid}").json()
        if r["status"] != "running" and r["crawl"]["status"] != "running": return r
        time.sleep(0.05)
    raise AssertionError("timed out")


def fake_engine(main, monkeypatch, fail=()):
    crawled = []

    class Done:
        def __init__(self, url, ok): self.url, self.ok = url, ok
        def jsonable(self):
            host = self.url.split("//")[1].strip("/")
            if not self.ok: return {"status": "failed", "stats": {"error": "blocked"}, "enrichment": None}
            return {"status": "completed", "stats": {}, "enrichment": {"status": "done", "lead_score": {"score": 71, "grade": "B"},
                    "contacts": {"emails": [{"email": "info@" + host}], "phones": [{"raw": "021 345", "e164": "+922134567890", "valid": True}]},
                    "people": [{"name": "Sara Khan", "title": "CEO", "email": None, "email_guess": "sara@" + host}],
                    "social": {"linkedin": {"url": "https://linkedin.com/company/x"}}}}

    async def run(job_id, url, cfg):
        crawled.append((url, cfg["max_pages"], cfg["registry_lookup"])); return Done(url, not any(f in url for f in fail))

    async def public(url): return True
    monkeypatch.setattr(main.ENGINE, "run", run); monkeypatch.setattr(main.security, "is_public_url", public); monkeypatch.setattr(main, "domain_resolves", public)
    return crawled


def test_search_auto_crawls_every_company_with_a_website_and_exports_one_sheet_each(client, monkeypatch):
    import main
    crawled = fake_engine(main, monkeypatch)
    rid = client.post("/api/companies/search", json={"industry": "textile", "country": "pk", "limit": 25}).json()["id"]
    r = wait(client, rid)
    cs = r["companies"]; sites = [c for c in cs if c.get("website")]
    assert len(cs) == 25 and r["total_count"] > 300 and len(crawled) == len(sites) == r["crawl"]["total"] and r["crawl"]["done"] == len(sites)
    assert all(__import__("urllib.parse").parse.urlparse(u).path in ("", "/") for u, _, _ in crawled)              # home page only, never a deep link
    assert all(p == 6 and reg is False for _, p, reg in crawled)
    c0 = cs[0]
    assert c0["crawl_status"] == "completed" and c0["details"]["decision_maker"] == "Sara Khan" and c0["profile"]["phones"][0]["e164"] == "+922134567890"
    x = client.get(f"/api/companies/{rid}/export/xlsx")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Sources" and len(wb.sheetnames) == 25 + 2
    assert len(set(n.lower() for n in wb.sheetnames)) == len(wb.sheetnames) and all(len(n) <= 31 for n in wb.sheetnames)
    ws = wb["Summary"]; header = [c.value for c in ws[1]]
    assert ws.max_row == 26 and {"name", "website", "all_emails", "decision_maker", "lead_score"} <= set(header)
    one = wb[wb.sheetnames[1]]
    flat = {row[0].value: row[1].value for row in one.iter_rows() if row[0].value}
    assert flat["Company"] == c0["name"] and "People" in flat and flat["lead score"] == 71
    csv = client.get(f"/api/companies/{rid}/export/csv").text.splitlines()
    assert len(csv) == 26 and "name" in csv[0]
    assert client.get(f"/api/companies/{rid}/export/pdf").status_code == 400
    assert client.get("/api/companies").json()[0]["count"] == 25


def test_a_failed_crawl_does_not_stop_the_others_and_is_reported(client, monkeypatch):
    import main
    crawled = fake_engine(main, monkeypatch, fail=("."))    # every site fails
    rid = client.post("/api/companies/search", json={"industry": "textile", "country": "pk", "limit": 5}).json()["id"]
    r = wait(client, rid)
    assert r["crawl"]["done"] == r["crawl"]["total"] and all(c.get("crawl_status") == "failed" for c in r["companies"] if c.get("website"))
    wb = load_workbook(io.BytesIO(client.get(f"/api/companies/{rid}/export/xlsx").content))
    assert len(wb.sheetnames) == 7                                   # company rows still exported with their listed contacts


def test_validation(client):
    assert client.post("/api/companies/search", json={"industry": "", "country": ""}).status_code == 400
    assert client.post("/api/companies/search", json={"industry": "textile", "country": ""}).status_code == 400
    rid = client.post("/api/companies/search", json={"industry": "textile", "country": "Nowhereland"}).json()["id"]
    r = wait(client, rid)
    assert r["status"] == "failed" and "Unknown country" in r["error"]
    assert client.get("/api/companies/nope").status_code == 404


def test_existing_endpoints_and_login(client, monkeypatch):
    assert client.get("/").status_code == 200
    assert client.get("/api/history").json() == [] and client.get("/api/profiles").json() == []
    assert client.post("/api/jobs", json={"url": "http://127.0.0.1/"}).status_code == 400
    monkeypatch.setenv("APP_PASSWORD", "pw")
    assert client.get("/api/companies").status_code == 401
